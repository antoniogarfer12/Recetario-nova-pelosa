"""Genera recetas.xlsx a partir de recetas.json."""
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
recetas = json.loads((RAIZ / "recetas.json").read_text(encoding="utf-8"))
recetas.sort(key=lambda r: (r.get("cat") or "", (r.get("name") or "").lower()))

wb = Workbook()
ws = wb.active
ws.title = "Recetas"
columnas = [
    ("Nombre", 28), ("Categoría", 14), ("Hecha por", 16), ("Tiempo (min)", 13), ("Raciones", 10),
    ("Ingredientes", 40), ("Pasos", 60), ("Notas", 35), ("Favorita", 10),
    ("Última modificación", 20),
]
ws.append([c for c, _ in columnas])
for i, (_, ancho) in enumerate(columnas, start=1):
    ws.column_dimensions[get_column_letter(i)].width = ancho
for celda in ws[1]:
    celda.font = Font(bold=True)
    celda.fill = PatternFill("solid", fgColor="E2B33C")

for r in recetas:
    pasos = r.get("steps") or []
    actualizado = r.get("updated")
    ws.append([
        r.get("name", ""),
        r.get("cat", ""),
        r.get("author", ""),
        r.get("time"),
        r.get("serv"),
        "\n".join(r.get("ing") or []),
        "\n".join(f"{n}. {p}" for n, p in enumerate(pasos, start=1)),
        r.get("notes", ""),
        "Sí" if r.get("fav") else "No",
        datetime.fromtimestamp(actualizado / 1000, tz=timezone.utc).replace(tzinfo=None) if actualizado else None,
    ])

for fila in ws.iter_rows(min_row=2):
    for celda in fila:
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    fila[9].number_format = "dd/mm/yyyy hh:mm"

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
wb.save(RAIZ / "recetas.xlsx")
print(f"recetas.xlsx generado con {len(recetas)} recetas")
